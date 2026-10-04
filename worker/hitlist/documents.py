"""Project documents from Dropbox: find each project's folder, read the useful sub-folders page by page,
and hand the text to the website so AI Tools can answer questions about specs and submittals.

Every project folder has the same sub-folders; only these are read (in this order):
Drawings and Specs, Submittal, ASIs and RFIs, Change Orders, Deficiency Reports (every punch list sent).

What the website gets (all gzipped JSON, stored privately in Vercel Blob under docs/<project>/):
  manifest  {project, folder, updated_at, files: [{id, category, name, path, modified, hash, pages, parts}], skipped}
  <part>    {file_id, start, pages: [text of each page]}   (big files are split into parts of PART_PAGES pages)
Unchanged files (same Dropbox content hash) are not downloaded or sent again.
"""

from __future__ import annotations

import gzip
import json
import logging
import re
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Callable

from .dropbox import Dropbox, DropboxError, find_folder
from .matching import similarity

log = logging.getLogger("hitlist")

DOC_FOLDERS = ["Drawings and Specs", "Submittal", "ASIs and RFIs", "Change Orders", "Deficiency Reports"]
READABLE = (".pdf", ".docx", ".xlsx", ".xlsm", ".txt")
SKIP_TECH_FOLDERS = re.compile(r"^\d\d\s")          # "01 Tech Resources" and the like are not technicians
PART_PAGES = 300                                      # pages per uploaded part (keeps each upload small)
MAX_PAGE_CHARS = 20_000
MAX_FILE_MB = 400


# Sub-folders holding superseded copies (an old spec next to the current one) aren't read, so answers only
# come from current documents: "Older versions", "Old", "Superseded", "Archive", "Previous", "Void"...
OLD_FOLDER = re.compile(r"^\W*(old(er)?( versions?| revisions?| docs?| files?)?|superseded|archived?|obsolete|"
                        r"previous( versions?| revisions?)?|prior( versions?)?|void(ed)?|outdated)\W*$", re.I)


def is_old_copy(rel_path: str) -> bool:
    """True when any folder in the file's path (inside the document folder) is an old-versions folder."""
    return any(OLD_FOLDER.match(part) for part in rel_path.split("/")[:-1])


class DocsError(RuntimeError):
    pass


# ---- finding the project's folder ---------------------------------------------------------
def starts_with_number(folder_name: str, number: str) -> bool:
    """'99-001 Some Job' starts with 99-001; '99-0012 Other' does not."""
    return bool(number) and re.match(rf"^{re.escape(number.strip())}(?![0-9A-Za-z])", folder_name.strip(), re.I) is not None


def _words(text: str) -> set[str]:
    return set(re.findall(r"[a-z0-9]+", (text or "").lower()))


def best_by_name(items: list[dict], name: str, label=lambda f: f["name"]) -> dict | None:
    """The item whose name clearly matches `name`, or None when it isn't clear. Words every item shares
    (the job name, "ES", "Elementary"...) are ignored, so "North Elementary" vs "South Elementary" is decided by
    North/South. Falls back to overall similarity when no distinguishing word appears in the name."""
    want = _words(name)
    common = set.intersection(*(_words(label(f)) for f in items)) if items else set()
    def score(f: dict) -> float:
        own = _words(label(f)) - common
        return len(own & want) / len(own) if own else 0.0
    ranked = sorted(items, key=score, reverse=True)
    if len(ranked) == 1:
        return ranked[0]
    if score(ranked[0]) > 0 and score(ranked[0]) > score(ranked[1]):
        return ranked[0]
    ranked = sorted(items, key=lambda f: similarity(label(f), name or ""), reverse=True)
    best, second = similarity(label(ranked[0]), name or ""), similarity(label(ranked[1]), name or "")
    return ranked[0] if best >= 0.35 and best - second >= 0.1 else None


def folder_index(dbx: Dropbox, tech_root: str) -> list[dict]:
    """Every folder inside the technician folders, plus folders sitting directly under the technicians folder:
    the places a project folder can be. Listing this once lets many projects be matched cheaply."""
    top = find_folder(dbx.folders(""), tech_root)
    if not top:
        raise DocsError(f"Couldn't find the '{tech_root}' folder in Dropbox")
    out: list[dict] = []
    for f in dbx.folders(top["path_lower"]):
        out.append(f)
        if not SKIP_TECH_FOLDERS.match(f["name"]) and not re.match(r"^\d{2}-\d", f["name"]):
            out += dbx.folders(f["path_lower"])
    return out


def _tech_of(folder: dict) -> str:
    """The technician folder a project folder sits in ("" when it sits directly under the technicians folder)."""
    parts = [p for p in folder.get("path_display", "").split("/") if p]
    return parts[-2] if len(parts) >= 3 else ""


def _where(folder: dict) -> str:
    return f"{_tech_of(folder)} / {folder['name']}" if _tech_of(folder) else folder["name"]


def _techs_match(folder_tech: str, tech: str) -> bool:
    """Hitlist's tech ("Sam" or "Sam Tech") against a technician folder name ("Sam Tech")."""
    a, b = _words(folder_tech), _words(tech)
    return bool(a and b) and (b <= a or a <= b)


def pick_project_folder(index: list[dict], number: str, name: str, taken: set[str] | None = None,
                        tech: str = "") -> dict | None:
    """The project's folder: the one starting with the job number. With more than one, the folder in the
    project's tech's own folder wins (copies of a job often sit in several techs' folders); then the Hitlist
    name decides (several projects can share a job number), but only when one clearly matches best; otherwise
    it stops and lists them so someone picks with Set folder. Folders already linked to another project with
    this number are skipped; the same folder seen twice counts once."""
    seen: dict[str, dict] = {}
    for f in index:
        if starts_with_number(f["name"], number):
            seen.setdefault(f["id"], f)
    candidates = list(seen.values())
    free = [c for c in candidates if c["id"] not in (taken or set())]
    if candidates and not free:
        raise DocsError(f"The only folder starting with {number} ({_where(candidates[0])}) is already linked to the "
                        f"other project with this number. If they share it, choose it with Set folder on the AI Tools tab.")
    if not free:
        return None
    if len(free) == 1:
        return free[0]
    mine = [f for f in free if tech and _techs_match(_tech_of(f), tech)]
    if len(mine) == 1:
        return mine[0]
    pool = mine or free
    rest = lambda f: re.sub(rf"^{re.escape(number)}\W*", "", f["name"], flags=re.I)  # noqa: E731
    best = best_by_name(pool, name, rest)
    if best is None:
        names = "; ".join(_where(f) for f in pool[:6])
        raise DocsError(f"Several Dropbox folders start with {number} and Hitlist can't tell which is this project's "
                        f"({names}). Choose it with Set folder on the AI Tools tab.")
    return best


def find_project_folder(dbx: Dropbox, tech_root: str, number: str, name: str,
                        taken: set[str] | None = None, tech: str = "") -> dict | None:
    return pick_project_folder(folder_index(dbx, tech_root), number, name, taken, tech)


# Folders every project folder has; two or more of them means "this is the project's document folder".
STANDARD_FOLDERS = DOC_FOLDERS + ["TAB Plan", "Contract", "Field Notes", "Final Report"]


def folder_key(name: str) -> str:
    """A folder name reduced so small differences don't matter: "02 - Submittals", "Drawings & Specs",
    "ASI's and RFI's" match "Submittal", "Drawings and Specs", "ASIs and RFIs"."""
    t = name.lower().replace("&", " and ").replace("'", "")
    t = re.sub(r"^[\s\d._-]+", "", t)                         # a leading number
    words = [w[:-1] if len(w) > 3 and w.endswith("s") else w for w in re.findall(r"[a-z0-9]+", t)]
    return " ".join(words)


def match_folder(folders: list[dict], wanted: str) -> dict | None:
    """The folder that is `wanted` (exact name, or the same apart from small differences)."""
    key = folder_key(wanted)
    exact = [f for f in folders if folder_key(f["name"]) == key]
    if exact:
        return exact[0]
    words = set(key.split())
    loose = [f for f in folders if words <= set(folder_key(f["name"]).split())]
    return loose[0] if len(loose) == 1 else None


def _is_project_level(dbx: Dropbox, folder: dict) -> bool:
    subs = dbx.folders(folder["path_lower"])
    return sum(match_folder(subs, s) is not None for s in STANDARD_FOLDERS) >= 2


def documents_folder(dbx: Dropbox, folder: dict, name: str) -> dict:
    """The folder that holds Drawings and Specs, Submittal, etc. Usually the job folder itself; when one job
    number covers several sites, the job folder holds a folder per site and the standard folders are inside
    those. Then the site folder whose name clearly matches the Hitlist name is used (one site: that one)."""
    if _is_project_level(dbx, folder):
        return folder
    sites = [f for f in dbx.folders(folder["path_lower"]) if _is_project_level(dbx, f)]
    if not sites:
        return folder                                   # no standard folders anywhere yet: keep the job folder
    if len(sites) == 1:
        return sites[0]
    best = best_by_name(sites, name)
    if best is None:
        raise DocsError(f"{folder['name']} has a folder for each site ({'; '.join(f['name'] for f in sites[:6])}) and the "
                        f"project name doesn't clearly pick one. Choose it with Set folder on the AI Tools tab.")
    return best


# ---- last punch list sent -----------------------------------------------------------------
PUNCH_FOLDER = "Deficiency Reports"


def last_punch_list(dbx: Dropbox, folder: dict) -> dict | None:
    """The newest file in the project's Deficiency Reports folder (techs save each sent punch list there):
    {"name", "date"}; {} when the folder is empty; None when the project folder has no such folder.
    Only names and dates are read, never the files."""
    sub = match_folder(dbx.folders(folder["path_lower"]), PUNCH_FOLDER)
    if not sub:
        return None
    files = [e for e in dbx.list_all(sub["path_lower"]) if e.get(".tag") == "file"]
    if not files:
        return {}
    when = lambda e: e.get("client_modified") or e.get("server_modified") or ""  # noqa: E731
    newest = max(files, key=when)
    return {"name": newest["name"], "date": when(newest)}


def dropbox_path_from(text: str) -> str:
    """Turn what someone pasted into a Dropbox path: a Windows path from the synced folder
    (C:\\Users\\...\\Company Dropbox\\PFE - Technician\\...), a dropbox.com/home/... link, or a plain path."""
    from urllib.parse import unquote, urlparse
    raw = text.strip().strip('"').strip()
    if re.match(r"^https?://", raw, re.I):
        path = re.sub(r"^/home(?=/|$)", "", unquote(urlparse(raw).path))
        return path.rstrip("/") or "/"
    parts = [p for p in raw.replace("\\", "/").split("/") if p]
    for i, p in enumerate(parts):                      # drop everything up to the synced "... Dropbox" folder
        if re.search(r"dropbox( \(.*\))?$", p, re.I):
            parts = parts[i + 1:]
            break
    if parts and re.match(r"^[a-z]:$", parts[0], re.I):
        raise DocsError("That path isn't inside the company Dropbox folder. Copy the folder's path from File Explorer "
                        "inside the Dropbox folder, or the link from the Dropbox website.")
    return "/" + "/".join(parts)


# ---- reading files ------------------------------------------------------------------------
def _clean(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text.replace("\x00", ""))
    return re.sub(r"\n{3,}", "\n\n", text).strip()[:MAX_PAGE_CHARS]


def read_pages(path: Path) -> list[str]:
    """Text of each page (PDF), each sheet (Excel) or each ~3000-character chunk (Word, text)."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pymupdf
        # Many CAD-exported PDFs have small internal glitches MuPDF repairs on the fly; don't print each one.
        pymupdf.TOOLS.mupdf_display_errors(False)
        pymupdf.TOOLS.mupdf_display_warnings(False)
        with pymupdf.open(path) as doc:
            return [_clean(page.get_text("text")) for page in doc]
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        pages = []
        for ws in wb.worksheets:
            rows = [" | ".join(str(v) for v in row if v not in (None, "")) for row in ws.iter_rows(values_only=True)]
            pages.append(_clean(f"[Sheet: {ws.title}]\n" + "\n".join(r for r in rows if r)))
        wb.close()
        return pages
    if ext == ".docx":
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
        xml = re.sub(r"</w:p>", "\n", xml)
        xml = re.sub(r"<w:tab/>", "\t", xml)
        text = re.sub(r"<[^>]+>", "", xml)
        return _chunks(text)
    if ext == ".txt":
        return _chunks(path.read_text(encoding="utf-8", errors="ignore"))
    return []


def _chunks(text: str, size: int = 3000) -> list[str]:
    text = _clean(text) if len(text) <= MAX_PAGE_CHARS else re.sub(r"[ \t]+", " ", text)
    out, cur = [], ""
    for para in text.split("\n"):
        if len(cur) + len(para) > size and cur:
            out.append(cur.strip())
            cur = ""
        cur += para + "\n"
    if cur.strip():
        out.append(cur.strip())
    return out or [""]


# ---- the whole update -----------------------------------------------------------------------
Upload = Callable[[str, bytes], None]      # (part name, gzipped JSON bytes)


def gz(obj: Any) -> bytes:
    return gzip.compress(json.dumps(obj, ensure_ascii=False).encode("utf-8"))


def update_documents(dbx: Dropbox, folder: dict, project: str, previous: dict | None, upload: Upload,
                     step: Callable[[str], None] = lambda m: None, now: str = "") -> dict:
    """Read the project's document folders and upload what changed. Returns the new manifest
    (also uploaded last, so the website never points at parts that aren't there yet)."""
    old = {f["id"]: f for f in (previous or {}).get("files", []) if f.get("parts")}
    subs = dbx.folders(folder["path_lower"])
    files: list[dict] = []
    skipped: list[dict] = []
    found: dict[str, int] = {}
    missing: list[str] = []                 # standard folders not in the project folder at all
    other: dict[str, list[str]] = {}        # files of types that aren't read (.dwg, .zip, images...), by folder
    for category in DOC_FOLDERS:
        sub = match_folder(subs, category)
        found[category] = 0
        if not sub:
            missing.append(category)
            continue
        for e in sorted((e for e in dbx.list_all(sub["path_lower"]) if e.get(".tag") == "file"),
                        key=lambda e: e["path_lower"]):
            rel = e["path_display"][len(sub["path_display"]):].lstrip("/")
            if is_old_copy(rel):
                continue
            if not e["name"].lower().endswith(READABLE):
                ext = Path(e["name"]).suffix.lower() or "(no type)"
                if ext not in other.setdefault(category, []):
                    other[category].append(ext)
                continue
            found[category] += 1
            entry = {"id": e["id"], "category": category, "name": e["name"], "path": rel,
                     "modified": e.get("server_modified", ""), "hash": e.get("content_hash", ""), "size": e.get("size", 0)}
            prev = old.get(e["id"])
            if prev and prev.get("hash") == entry["hash"]:
                files.append({**entry, "pages": prev["pages"], "parts": prev["parts"]})
                continue
            if entry["size"] > MAX_FILE_MB * 1024 * 1024:
                skipped.append({**entry, "reason": f"larger than {MAX_FILE_MB} MB"})
                continue
            step(f"Reading {category}: {e['name']}")
            try:
                with tempfile.TemporaryDirectory() as tmp:
                    pages = read_pages(dbx.download(e["id"], Path(tmp) / ("f" + Path(e["name"]).suffix.lower())))
            except DropboxError:
                raise
            except Exception as exc:  # noqa: BLE001 - a damaged or locked file shouldn't stop the rest
                log.warning("Could not read %s: %s", e["path_display"], exc)
                skipped.append({**entry, "reason": "couldn't be read (damaged or password-protected?)"})
                continue
            if not any(p.strip() for p in pages):
                skipped.append({**entry, "reason": "no text in it (a scan or image-only file)"})
                continue
            stem = re.sub(r"[^A-Za-z0-9]", "", e["id"])[-12:] + "-" + re.sub(r"[^A-Za-z0-9]", "", entry["hash"])[:10]
            parts = []
            for start in range(0, len(pages), PART_PAGES):
                name = f"{stem}-{start // PART_PAGES}"
                upload(name, gz({"file_id": e["id"], "start": start, "pages": pages[start:start + PART_PAGES]}))
                parts.append(name)
            files.append({**entry, "pages": len(pages), "parts": parts})
    manifest = {"project": project, "folder": {"id": folder["id"], "path": folder.get("path_display", ""),
                                               "name": folder.get("name", "")},
                "updated_at": now, "found": found, "missing": missing, "other_types": other,
                "files": files, "skipped": skipped}
    step("Saving document index")
    upload("manifest", gz(manifest))
    return manifest
